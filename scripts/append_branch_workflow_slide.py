from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PPTX = (
    ROOT
    / "Files"
    / "PPT-Producer"
    / "Teaching-Associate教授の助理 职位"
    / "GitHub三大核心概念.pptx"
)
IMAGE = (
    ROOT
    / "Files"
    / "PPT-Producer"
    / "Teaching-Associate教授の助理 职位"
    / "在未Merge条件下更新PPT文件.png"
)

FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
W, H = 1600, 853


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size)


def text(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    value: str,
    *,
    size: int = 22,
    color: str = "#1F2328",
) -> None:
    draw.text(xy, value, font=font(size), fill=color)


def rounded(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    *,
    fill: str,
    outline: str = "#D0D7DE",
    radius: int = 8,
    width: int = 2,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def create_guide_image() -> None:
    image = Image.new("RGB", (W, H), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    # GitHub-style top navigation.
    draw.rectangle((0, 0, W, 92), fill="#F6F8FA")
    draw.line((0, 91, W, 91), fill="#D8DEE4", width=2)
    rounded(draw, (14, 16, 54, 56), fill="#FFFFFF", radius=7)
    text(draw, (27, 18), "≡", size=25)
    draw.ellipse((71, 14, 113, 56), fill="#1F2328")
    text(draw, (78, 19), "GH", size=13, color="#FFFFFF")
    text(draw, (126, 20), "36613453226Wallen  /  R-Calendar", size=18)
    rounded(draw, (955, 13, 1192, 57), fill="#FFFFFF", radius=8)
    text(draw, (980, 22), "Type / to search", size=15, color="#57606A")
    for x, label in [
        (1230, "+"),
        (1285, "○"),
        (1340, "⑂"),
        (1395, "▣"),
        (1450, "⌂"),
    ]:
        rounded(draw, (x, 15, x + 42, 57), fill="#FFFFFF", radius=7)
        text(draw, (x + 13, 21), label, size=17, color="#57606A")
    draw.ellipse((1520, 14, 1562, 56), fill="#FFF3BF")

    nav = ["Code", "Issues", "Pull requests  2", "Agents", "Actions", "Projects", "Wiki"]
    x = 16
    for index, label in enumerate(nav):
        color = "#1F2328" if index == 0 else "#57606A"
        text(draw, (x, 67), label, size=14, color=color)
        x += 58 + len(label) * 6
    draw.line((10, 89, 72, 89), fill="#FD8C73", width=3)

    # Left file tree.
    draw.rectangle((0, 92, 378, H), fill="#F6F8FA")
    draw.line((377, 92, 377, H), fill="#D0D7DE", width=2)
    text(draw, (18, 115), "▣  Files", size=21)
    text(draw, (145, 162), "Step1 Branch", size=26, color="#FF4D4F")
    rounded(draw, (12, 188, 272, 232), fill="#FFFFFF", radius=7)
    text(draw, (25, 200), "⑂  cursor/create-g...", size=16, color="#57606A")
    rounded(draw, (278, 188, 317, 232), fill="#FFFFFF", radius=7)
    text(draw, (290, 199), "+", size=18, color="#57606A")
    rounded(draw, (323, 188, 363, 232), fill="#FFFFFF", radius=7)
    text(draw, (335, 200), "⌕", size=17, color="#57606A")
    rounded(draw, (12, 244, 363, 282), fill="#FFFFFF", radius=7)
    text(draw, (25, 252), "⌕  Go to file", size=15, color="#57606A")
    text(
        draw,
        (22, 303),
        "⌄  📁 Files/PPT-Producer/Teaching-Ass...",
        size=14,
        color="#0969DA",
    )
    draw.rectangle((8, 338, 370, 381), fill="#DDF4FF")
    text(draw, (42, 347), "▯  GitHub三大核心概念.pptx", size=15)
    text(draw, (317, 347), "Step2", size=25, color="#FF4D4F")
    text(draw, (294, 382), "选择Branch下的文件", size=22, color="#FF4D4F")
    text(draw, (43, 423), "▯  Teaching-Associate职位对照表...", size=14)
    text(draw, (21, 469), "›  📁 scripts", size=15, color="#0969DA")
    text(draw, (42, 510), "▯  .gitignore", size=15)
    text(draw, (42, 550), "▯  README.md", size=15)

    # Main file page.
    text(
        draw,
        (402, 111),
        "R-Calendar  /  Files  /  PPT-Producer  /  Teaching-Associate教授の助理 职位  /",
        size=16,
        color="#0969DA",
    )
    text(draw, (1246, 111), "GitHub三大核心概念.pptx", size=16)
    rounded(draw, (403, 148, 1564, 212), fill="#FFFFFF", radius=8)
    draw.ellipse((422, 166, 449, 193), fill="#FFF3BF")
    text(draw, (463, 169), "36613453226Wallen", size=14)
    text(draw, (622, 169), "Add files via upload", size=14, color="#57606A")
    rounded(draw, (402, 229, 1564, 690), fill="#FFFFFF", radius=7)
    rounded(draw, (416, 242, 478, 279), fill="#F6F8FA", radius=6)
    text(draw, (430, 250), "Code", size=14)
    text(draw, (493, 250), "Blame", size=14, color="#57606A")
    text(draw, (546, 252), "61.9 KB", size=13, color="#8C959F")
    text(draw, (966, 291), "View raw", size=14, color="#0969DA")

    # Ellipsis menu.
    rounded(draw, (1550, 104, 1592, 145), fill="#F6F8FA", radius=8)
    text(draw, (1561, 106), "•••", size=17, color="#57606A")
    rounded(draw, (1282, 145, 1584, 698), fill="#FFFFFF", radius=12)
    menu = [
        ("Raw file content", 168, "#1F2328"),
        ("Download", 207, "#1F2328"),
        ("Copy path", 266, "#1F2328"),
        ("Copy permalink", 305, "#1F2328"),
        ("Copilot", 368, "#1F2328"),
        ("Ask about this file", 407, "#1F2328"),
        ("View options", 469, "#57606A"),
        ("✓  Show code folding buttons", 509, "#1F2328"),
        ("   Wrap lines", 548, "#57606A"),
        ("   Center content", 587, "#57606A"),
        ("✓  Open symbols on click", 626, "#1F2328"),
    ]
    for label, y, color in menu:
        text(draw, (1302, y), label, size=14, color=color)
    draw.line((1295, 243, 1570, 243), fill="#D8DEE4", width=1)
    draw.line((1295, 346, 1570, 346), fill="#D8DEE4", width=1)
    draw.line((1295, 447, 1570, 447), fill="#D8DEE4", width=1)
    draw.rectangle((1288, 653, 1578, 693), fill="#FFEBE9")
    text(draw, (1302, 662), "Delete file", size=14, color="#CF222E")

    # Instructional annotations from the supplied screenshot.
    text(draw, (1044, 699), "Step3 选择删除文件", size=26, color="#FF4D4F")
    text(draw, (958, 741), "Commit changes 同意更改", size=26, color="#2ECC71")
    text(
        draw,
        (958, 785),
        "“在未Merge条件下进行更改の操作可行”",
        size=25,
        color="#1F2328",
    )

    image.save(IMAGE, format="PNG", dpi=(200, 200), optimize=True)


def append_slide() -> None:
    prs = Presentation(PPTX)
    if len(prs.slides) != 5:
        raise RuntimeError(
            f"Expected the user-updated five-slide deck, found {len(prs.slides)} slides"
        )
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(IMAGE),
        0,
        0,
        width=prs.slide_width,
        height=prs.slide_height,
    )
    with NamedTemporaryFile(
        suffix=".pptx", dir=PPTX.parent, delete=False
    ) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        prs.save(temp_path)
        temp_path.replace(PPTX)
    finally:
        temp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    create_guide_image()
    append_slide()
    print(PPTX)
    print(IMAGE)
