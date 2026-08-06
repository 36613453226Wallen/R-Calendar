from __future__ import annotations

import calendar
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "document"
PNG_PATH = OUTPUT_DIR / "2026年8-9月排练演出日历.png"
DOCX_PATH = OUTPUT_DIR / "2026年8-9月排练演出日历.docx"

WIDTH, HEIGHT = 2480, 1754
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"

BACKGROUND = "#F6F8FC"
CARD = "#FFFFFF"
TEXT = "#253248"
MUTED = "#738096"
GRID = "#D9E1EC"
WEEKEND = "#FAFBFD"
REHEARSAL = "#1A73E8"
REHEARSAL_LIGHT = "#E8F0FE"
PERFORMANCE = "#0F9D58"
PERFORMANCE_LIGHT = "#E6F4EA"
HOLIDAY = "#C5221F"

WEEKDAYS = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
REHEARSALS = {
    (2026, 8, 16),
    (2026, 8, 23),
    (2026, 8, 30),
    (2026, 9, 6),
    (2026, 9, 13),
    (2026, 9, 20),
}
PERFORMANCES = {
    (2026, 9, 25),
    (2026, 9, 26),
    (2026, 9, 27),
}


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_PATH, size)


def center_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    text_font: ImageFont.FreeTypeFont,
    fill: str,
) -> None:
    left, top, right, bottom = box
    bounds = draw.textbbox((0, 0), text, font=text_font)
    text_width = bounds[2] - bounds[0]
    text_height = bounds[3] - bounds[1]
    draw.text(
        (
            left + (right - left - text_width) / 2,
            top + (bottom - top - text_height) / 2 - bounds[1],
        ),
        text,
        font=text_font,
        fill=fill,
    )


def draw_event(
    draw: ImageDraw.ImageDraw,
    cell: tuple[int, int, int, int],
    label: str,
    color: str,
    light_color: str,
) -> None:
    left, top, right, _ = cell
    bar_top = top + 70
    bar_bottom = bar_top + 47
    draw.rounded_rectangle(
        (left + 12, bar_top, right - 12, bar_bottom),
        radius=12,
        fill=light_color,
    )
    draw.rounded_rectangle(
        (left + 12, bar_top, left + 20, bar_bottom),
        radius=4,
        fill=color,
    )
    center_text(
        draw,
        (left + 23, bar_top, right - 15, bar_bottom),
        label,
        font(25),
        color,
    )


def draw_month(
    draw: ImageDraw.ImageDraw,
    year: int,
    month: int,
    panel: tuple[int, int, int, int],
) -> None:
    left, top, right, bottom = panel
    draw.rounded_rectangle(panel, radius=28, fill=CARD, outline=GRID, width=2)

    month_header_height = 95
    weekday_height = 58
    grid_top = top + month_header_height + weekday_height
    col_width = (right - left) / 7
    row_height = (bottom - grid_top) / 6

    center_text(
        draw,
        (left, top + 5, right, top + month_header_height),
        f"{year}年{month}月",
        font(44),
        TEXT,
    )

    for col, weekday in enumerate(WEEKDAYS):
        x0 = round(left + col * col_width)
        x1 = round(left + (col + 1) * col_width)
        fill = HOLIDAY if col >= 5 else MUTED
        center_text(
            draw,
            (x0, top + month_header_height, x1, grid_top),
            weekday,
            font(24),
            fill,
        )

    weeks = calendar.Calendar(firstweekday=calendar.MONDAY).monthdayscalendar(
        year, month
    )
    while len(weeks) < 6:
        weeks.append([0] * 7)

    for row in range(6):
        for col in range(7):
            x0 = round(left + col * col_width)
            x1 = round(left + (col + 1) * col_width)
            y0 = round(grid_top + row * row_height)
            y1 = round(grid_top + (row + 1) * row_height)
            if col >= 5:
                draw.rectangle((x0, y0, x1, y1), fill=WEEKEND)
            draw.rectangle((x0, y0, x1, y1), outline=GRID, width=2)

            day = weeks[row][col]
            if not day:
                continue

            date_color = HOLIDAY if col >= 5 else TEXT
            draw.text(
                (x0 + 15, y0 + 14),
                str(day),
                font=font(30),
                fill=date_color,
            )

            key = (year, month, day)
            if key == (2026, 9, 25):
                holiday_bounds = draw.textbbox((0, 0), "中秋节", font=font(21))
                holiday_width = holiday_bounds[2] - holiday_bounds[0]
                draw.text(
                    (x1 - holiday_width - 13, y0 + 18),
                    "中秋节",
                    font=font(21),
                    fill=HOLIDAY,
                )
            if key in REHEARSALS:
                draw_event(
                    draw,
                    (x0, y0, x1, y1),
                    "排练",
                    REHEARSAL,
                    REHEARSAL_LIGHT,
                )
            if key in PERFORMANCES:
                draw_event(
                    draw,
                    (x0, y0, x1, y1),
                    "星光剧场演出",
                    PERFORMANCE,
                    PERFORMANCE_LIGHT,
                )


def create_image() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (WIDTH, HEIGHT), BACKGROUND)
    draw = ImageDraw.Draw(image)

    draw.text(
        (110, 80),
        "2026年8—9月",
        font=font(58),
        fill=TEXT,
    )
    draw.text(
        (110, 153),
        "排练与演出日历",
        font=font(34),
        fill=MUTED,
    )

    legend_y = 119
    rehearsal_x = 1740
    performance_x = 2025
    draw.rounded_rectangle(
        (rehearsal_x, legend_y, rehearsal_x + 28, legend_y + 28),
        radius=7,
        fill=REHEARSAL,
    )
    draw.text(
        (rehearsal_x + 42, legend_y - 3),
        "排练",
        font=font(27),
        fill=TEXT,
    )
    draw.rounded_rectangle(
        (performance_x, legend_y, performance_x + 28, legend_y + 28),
        radius=7,
        fill=PERFORMANCE,
    )
    draw.text(
        (performance_x + 42, legend_y - 3),
        "星光剧场演出",
        font=font(27),
        fill=TEXT,
    )

    draw_month(draw, 2026, 8, (80, 250, 1215, 1575))
    draw_month(draw, 2026, 9, (1265, 250, 2400, 1575))

    draw.text(
        (84, 1623),
        "排练：8月16、23、30日；9月6、13、20日",
        font=font(25),
        fill=MUTED,
    )
    performance_note = "演出：9月25—27日 · 星光剧场"
    note_bounds = draw.textbbox((0, 0), performance_note, font=font(25))
    draw.text(
        (2396 - (note_bounds[2] - note_bounds[0]), 1623),
        performance_note,
        font=font(25),
        fill=MUTED,
    )

    image.save(PNG_PATH, format="PNG", dpi=(300, 300), optimize=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_run_font(run, name: str, size: int, color: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = None
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    run._element.get_or_add_rPr().append(color_element)


def create_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("2026年8—9月排练与演出日历")
    title_run.bold = True
    set_run_font(title_run, "微软雅黑", 20, "253248")

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_after = Pt(5)
    image_paragraph.add_run().add_picture(str(PNG_PATH), width=Cm(24.0))

    summary = document.add_table(rows=1, cols=2)
    summary.autofit = False
    summary.columns[0].width = Cm(13.4)
    summary.columns[1].width = Cm(13.4)
    left_cell, right_cell = summary.rows[0].cells
    set_cell_shading(left_cell, "E8F0FE")
    set_cell_shading(right_cell, "E6F4EA")

    left_paragraph = left_cell.paragraphs[0]
    left_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left_run = left_paragraph.add_run("排练：8月16、23、30日；9月6、13、20日")
    set_run_font(left_run, "微软雅黑", 9, "1A73E8")

    right_paragraph = right_cell.paragraphs[0]
    right_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_run = right_paragraph.add_run("星光剧场演出：9月25—27日（中秋节档期）")
    set_run_font(right_run, "微软雅黑", 9, "0F9D58")

    document.save(DOCX_PATH)


if __name__ == "__main__":
    create_image()
    create_docx()
    print(PNG_PATH)
    print(DOCX_PATH)
